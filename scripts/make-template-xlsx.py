#!/usr/bin/env python3
"""Build templates/programme-template.xlsx from templates/programme-template.csv.

The Excel template has a Read me sheet, the Programme sheet (headings, sample
rows, drop-down lists for Fleet, Race start order, Tide source and GP weekend)
and a Lists sheet. High water and Start time are text cells, so they save to
CSV exactly as typed (14:52, TBC) whatever the spreadsheet's time settings.

    pip install openpyxl
    python3 scripts/make-template-xlsx.py templates/programme-template.csv templates/programme-template.xlsx
"""
import csv, datetime as dt, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
src, out = sys.argv[1], sys.argv[2]
rows = list(csv.DictReader(open(src, encoding="utf-8-sig")))
cols = ["Date","High water","Tide height (m)","Tide source","Start time","Fleet","Race start order","Event","Training","GP weekend","Race / series","Race no(s)","Detail"]
wb = Workbook()
ins = wb.active; ins.title = "Read me"
head = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="0E2231")
lines = [
 ("BSC Sailing Calendar: season programme template", True),
 ("", False),
 ("1. Fill in the Programme sheet: one row per day per fleet (or per event or training session). Delete the sample rows first.", False),
 ("2. Keep the headings in row 1 exactly as they are. Extra columns are fine; the app ignores them.", False),
 ("3. Dates are year-month-day (2027-04-18). Times are 24-hour (14:52). Start time can also be TBC.", False),
 ("4. Fleet, Race start order, Tide source and GP weekend have drop-down lists. Fleet names must match exactly.", False),
 ("5. Save the Programme sheet as CSV: File > Save As > CSV UTF-8, named programme-2027.csv. Excel warns only the current sheet is saved; that's expected.", False),
 ("6. Check it before sending: https://bsc-sailing.github.io/calendar/?preview (in the app) and https://bsc-sailing.github.io/calendar/racecard.html (as a race card). Nothing is uploaded.", False),
 ("", False),
 ("The full guide, with a worked example: https://bsc-sailing.github.io/calendar/guide.html", False),
]
for i,(t,b) in enumerate(lines,1):
    c = ins.cell(row=i, column=1, value=t); c.font = Font(bold=b, size=14 if b else 11)
ins.column_dimensions["A"].width = 130
ws = wb.create_sheet("Programme")
for j,h in enumerate(cols,1):
    c = ws.cell(row=1, column=j, value=h); c.font = head; c.fill = fill; c.alignment = Alignment(vertical="center")
def tm(s):
    return s or None   # times are kept as text, so they save exactly as typed (14:52, TBC)
for i,r in enumerate(rows,2):
    vals = [dt.datetime.strptime(r["Date"],"%Y-%m-%d"), tm(r["High water"]), float(r["Tide height (m)"]) if r["Tide height (m)"] else None,
            r["Tide source"] or None, tm(r["Start time"]), r["Fleet"] or None, int(r["Race start order"]) if r["Race start order"] else None,
            r["Event"] or None, r["Training"] or None, r["GP weekend"].upper()=="TRUE", r["Race / series"] or None, r["Race no(s)"] or None, r["Detail"] or None]
    for j,v in enumerate(vals,1): ws.cell(row=i, column=j, value=v)
N = 2000
fmts = {1:"yyyy-mm-dd", 2:"@", 3:"0.0", 5:"@", 12:"@"}
for col,f in fmts.items():
    for i in range(2, N+1): ws.cell(row=i, column=col).number_format = f
widths = [12,11,14,12,11,16,15,22,18,11,26,11,34]
for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "B2"
ls = wb.create_sheet("Lists")
lists = {"A":["Fleet","Fast / Fireball","Medium","Wayfarer","Sprite","Short course","Fridays","Mirror","Cruiser"],
         "B":["Race start order",1,2,3,4,5,6],"C":["Tide source","Programme","Estimated"],"D":["GP weekend","TRUE","FALSE"],
         "E":["Training","Cadet training","Cadet coaching","Adult RYA","Wayfarer training"],
         "F":["Event (examples)","No racing","Evening race","Regatta","Club week","Cadet week","Open meeting","Cruiser weekend","Mirror sailing","Beach Club","Informal race"]}
for col,vals in lists.items():
    for i,v in enumerate(vals,1):
        c = ls[f"{col}{i}"]; c.value = v
        if i==1: c.font = Font(bold=True)
    ls.column_dimensions[col].width = 20
ls.cell(row=13, column=1, value="Fleet and Training lists can be extended: add a name at the bottom of the column. Event is free text; these are just the usual ones.")
def dv(rng, src, allow_other=False, msg=""):
    d = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=not allow_other)
    d.error = msg; d.errorTitle = "Not in the list"; ws.add_data_validation(d); d.add(rng)
dv(f"F2:F{N}", "=Lists!$A$2:$A$20", msg="Fleet names must match the list exactly (add new fleets on the Lists sheet).")
dv(f"G2:G{N}", "=Lists!$B$2:$B$7", msg="1 to 6, the order the fleet starts in.")
dv(f"D2:D{N}", "=Lists!$C$2:$C$3", msg="Programme, or Estimated.")
dv(f"J2:J{N}", "=Lists!$D$2:$D$3", msg="TRUE or FALSE.")
dv(f"I2:I{N}", "=Lists!$E$2:$E$20", allow_other=True)
dv(f"H2:H{N}", "=Lists!$F$2:$F$11", allow_other=True)
wb.active = 1
wb.save(out)
