#!/usr/bin/env python3
"""Convert a race-card spreadsheet (one row per day, one column per fleet)
into the app's programme format (one row per day per fleet or event).

Written for the 2027 race card framework, whose "programme" sheet has, from
row 3 on: Day, Date, two high waters with heights, the race tide and height
(columns H and I), two start times (J and K, worked out from the tide), the
six fleet columns in start order (L to Q) and Other events (R). The draft it
writes is a starting point to edit, not a finished programme: check the
notes it prints.

    pip install openpyxl
    python3 scripts/convert-race-card.py "Race card 27 framework v2.xlsx" programme-2027.csv

An .xls file needs saving as .xlsx first (Excel: File > Save As > Excel
Workbook; or LibreOffice: soffice --headless --convert-to xlsx file.xls).

What it does with each kind of cell:
- A name in a fleet column becomes that fleet's race. "Club Week" or "Cadet
  week" across the fleets becomes one row per fleet with that Event. The
  combined "Short/Friday/Mirror" column becomes Short course, Fridays and
  Mirror, all starting 5th, as in 2026.
- Holiday labels (GOOD FRIDAY, BANK HOLIDAY ...) become a club-wide Event.
- "NO RACE" is the sheet's own helper flag (start outside 09:30-16:00) and
  is left out: it marks every such day of the year, not decisions.
- Other events: coaching and training become Training rows with start TBC;
  open meetings, regattas and cadet racing become Events; anything it
  doesn't recognise is kept as an Event under its own name.
- Days with nothing in any fleet or event column are left out.
- High water and height come from the race tide columns; the start time from
  the second (tide-adjusted) start column. Both are marked Programme.
"""
import csv
import datetime as dt
import re
import sys

import openpyxl

COLUMNS = ["Date", "High water", "Tide height (m)", "Tide source", "Start time", "Fleet", "Race start order",
           "Event", "Training", "GP weekend", "Race / series", "Race no(s)", "Detail"]
# Race-card fleet columns (L to Q), with the app's fleet names and start order.
FLEET_COLS = [(11, ["Fast / Fireball"], 1), (12, ["Medium"], 2), (13, ["Wayfarer"], 3), (14, ["Sprite"], 4),
              (15, ["Short course", "Fridays", "Mirror"], 5), (16, ["Cruiser"], 6)]
OTHER_COL = 17
WEEKS = {"club week": "Club week", "cadet week": "Cadet week"}
HOLIDAYS = {"good friday": "Good Friday", "easter monday": "Easter Monday", "bank holiday": "Bank holiday"}


def hhmm(v):
    if isinstance(v, (dt.datetime, dt.time)):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, (int, float)):            # an Excel time as a fraction of a day
        m = round((v % 1) * 1440)
        return f"{m // 60 % 24:02d}:{m % 60:02d}"
    s = str(v or "").strip()
    m = re.fullmatch(r"(\d{1,2})[:.](\d\d)(?::\d\d)?", s)
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else ""


def other_event(text):
    """Map an 'Other events' cell to (Event, Training, Race/series, start TBC?)."""
    t = text.strip()
    low = t.lower()
    if "coaching" in low:
        return "", "Cadet coaching", "", True
    if "training" in low:
        return "", ("Cadet training" if "cadet" in low or low == "training" else t), "", True
    if "frostbite" in low:
        return "Cadet racing", "", t, False
    if "regatta" in low or "reid scott" in low:
        return "Cadet regatta" if "cadet" in low or "reid scott" in low else "Regatta", "", t, False
    if "open" in low:
        return "Open meeting", "", t, False
    return t, "", "", False


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__.split("\n\n")[2])
    src, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["programme"] if "programme" in wb.sheetnames else wb.worksheets[0]
    rows, notes, tentative, holidays, skipped = [], [], [], [], 0
    for r in range(3, ws.max_row + 1):
        v = [c.value for c in ws[r]] + [None] * 20
        d = v[1]
        if not isinstance(d, (dt.datetime, dt.date)):
            continue
        iso = d.strftime("%Y-%m-%d")
        hw, ht = hhmm(v[7]), v[8]
        ht = f"{float(ht):.1f}" if isinstance(ht, (int, float)) else ""
        start = hhmm(v[10]) or hhmm(v[9])
        base = {"Date": iso, "High water": hw, "Tide height (m)": ht, "Tide source": "Programme" if hw else "",
                "GP weekend": "FALSE"}
        day = []
        for col, fleets, order in FLEET_COLS:
            cell = str(v[col] or "").strip()
            if not cell or cell.upper() == "NO RACE":
                continue
            low = cell.lower()
            if low in HOLIDAYS:
                if not any(x["Event"] == HOLIDAYS[low] for x in day):
                    day.append(dict(base, **{"Event": HOLIDAYS[low]}))
                    holidays.append(f"{iso} {HOLIDAYS[low]}")
                continue
            for f in fleets:
                row = dict(base, **{"Start time": start, "Fleet": f, "Race start order": str(order)})
                if low in WEEKS:
                    row["Event"] = WEEKS[low]
                else:
                    m = re.fullmatch(r"(.*?)\s+(\d+\s*(?:&|and|-)\s*\d+|\d+)", cell)
                    row["Race / series"], row["Race no(s)"] = (m.group(1), m.group(2).replace(" ", "")) if m else (cell, "")
                day.append(row)
        other = str(v[OTHER_COL] or "").strip()
        if other:
            for part in re.split(r"\s*/\s*", other):
                ev, tr, series, tbc = other_event(part)
                day.append(dict(base, **{"Event": ev, "Training": tr, "Race / series": series,
                                         "Start time": "TBC" if tbc else start}))
                if "?" in part:
                    tentative.append(f"{iso} {part}")
        if not day:
            skipped += 1
            continue
        if not hw:
            notes.append(f"{iso}: no race tide in the sheet, so the tide is blank (the app will estimate it)")
        rows.extend(day)

    with open(out, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS, lineterminator="\r\n")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in COLUMNS})
    days = sorted({r["Date"] for r in rows})
    print(f"Wrote {len(rows)} rows for {len(days)} days to {out} ({skipped} days with nothing on were left out).")
    if days:
        print(f"From {days[0]} to {days[-1]}.")
    if holidays:
        print("Holiday labels kept as club-wide events: " + "; ".join(holidays))
    if tentative:
        print("Marked with '?' in the sheet (kept, with the '?'): " + "; ".join(tentative))
    for n in notes:
        print(n)


if __name__ == "__main__":
    main()
